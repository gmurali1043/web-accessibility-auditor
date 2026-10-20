"""Exports a list of per-page report dicts (see report.py) to Markdown,
HTML, or XLSX. Results stay grouped by page in every format, with an
overall conformance score and a page-by-page score."""
import datetime
import html as _html

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

STATUS_LABEL = {"PASS": "Pass", "FAIL": "Fail", "REVIEW": "Manual Review Required"}
STATUS_COLORS = {"PASS": "1E7B34", "FAIL": "C0392B", "REVIEW": "8E44AD"}
STATUS_FILL = {"PASS": "E9F7EF", "FAIL": "FDEDEC", "REVIEW": "F4ECF7"}
TIER_LABEL = {"automated": "Automated", "partial": "Partially Automated", "manual": "Manual Review Required"}


def _overall(reports):
    pass_n = sum(r["summary"]["pass"] for r in reports)
    fail_n = sum(r["summary"]["fail"] for r in reports)
    review_n = sum(r["summary"]["review"] for r in reports)
    determinable = pass_n + fail_n
    score = round((pass_n / determinable) * 100, 1) if determinable else None
    return {"pass": pass_n, "fail": fail_n, "review": review_n, "determinable": determinable, "score_pct": score}


def _score_str(score):
    return f"{score}%" if score is not None else "n/a"


def _issue_location(issue):
    return issue.get("selector") or "(page-level)"


def _wcag_str(row):
    if row["id"] == "GEN-4-1":
        return "All axe-core WCAG 2.2 A/AA rules (see each issue)"
    wcag = row.get("wcag") or []
    if not wcag:
        return "Best practice (no direct WCAG criterion)"
    return "; ".join(f"{w['sc']} {w['name']} ({w['level']})" for w in wcag)


def _issue_wcag(issue, row):
    """Axe issues carry their own SCs; custom-check issues use the row's."""
    if issue.get("wcag"):
        return ", ".join(issue["wcag"])
    if issue.get("best_practice"):
        return "Best practice"
    return ", ".join(w["sc"] for w in row.get("wcag") or []) or "Best practice"


def _manual_steps(row):
    """Plain-text manual test script for a row."""
    parts = []
    if row.get("tool"):
        parts.append(f"Tool: {row['tool']}")
    if row.get("procedure"):
        parts.append(f"Steps: {row['procedure']}")
    if row.get("outcomes"):
        parts.append(f"Pass/fail: {row['outcomes']}")
    return "\n".join(parts)


# ============================================================ Markdown =====
def to_markdown(reports, title="Web Accessibility Report"):
    overall = _overall(reports)
    lines = [f"# {title}", "", f"_Generated {datetime.datetime.now():%Y-%m-%d %H:%M}_", ""]
    lines.append(f"**{len(reports)} page(s) audited. Overall conformance score: {_score_str(overall['score_pct'])} "
                 f"({overall['pass']} pass / {overall['fail']} fail out of {overall['determinable']} "
                 f"automatically-determinable checks; {overall['review']} flagged for manual review.)**")
    lines.append("")

    if len(reports) > 1:
        lines.append("| # | Page | Pass | Fail | Manual Review | Score |")
        lines.append("|---|------|------|------|----------------|-------|")
        for i, r in enumerate(reports, start=1):
            s = r["summary"]
            lines.append(f"| {i} | [{r['display_name']}]({r['url']}) | {s['pass']} | {s['fail']} | "
                          f"{s['review']} | {_score_str(s['score_pct'])} |")
        lines.append("")

    for r in reports:
        s = r["summary"]
        lines.append(f"## {r['display_name']}")
        lines.append("")
        lines.append(f"- **URL:** {r['url']}")
        if r.get("title"):
            lines.append(f"- **Page title:** {r['title']}")
        lines.append(f"- **Page conformance score:** {_score_str(s['score_pct'])}")
        lines.append(f"- **Results:** {s['pass']} Pass / {s['fail']} Fail / {s['review']} Manual review required "
                      f"(of {s['total']} test cases: {s['automated_tier']} automated, "
                      f"{s['partial_tier']} partially automated, {s['manual_tier']} manual-only)")
        if r.get("errors"):
            lines.append(f"- **Notes:** {'; '.join(r['errors'])}")
        lines.append("")

        for row in r["rows"]:
            lines.append(f"### {row['id']} — {row['title']}  [{STATUS_LABEL[row['status']]}]")
            lines.append("")
            lines.append(f"- **VGAR ID:** {row['id']} · **Category:** {row['category']}")
            lines.append(f"- **WCAG:** {_wcag_str(row)}")
            lines.append(f"- **Automation:** {TIER_LABEL[row['automation']]} — {row['method']}")
            if row["status"] == "REVIEW" and row.get("procedure"):
                lines.append("- **How to test manually:**")
                if row.get("tool"):
                    lines.append(f"  - **Tool:** {row['tool']}")
                lines.append(f"  - **Steps:** {row['procedure']}")
                if row.get("outcomes"):
                    lines.append(f"  - **Pass / fail:** {' '.join(row['outcomes'].split())}")
            if row["issues"]:
                lines.append("- **Findings:**")
                for issue in row["issues"]:
                    lines.append(f"  - **Location:** `{_issue_location(issue)}`")
                    lines.append(f"    - **VGAR / WCAG:** {row['id']} / {_issue_wcag(issue, row)}")
                    lines.append(f"    - **Issue:** {issue['message']}")
                    if issue.get("html"):
                        lines.append(f"    - **Code:** `{issue['html'][:200]}`")
                    if issue.get("fix"):
                        lines.append(f"    - **Fix:** {issue['fix']}")
            lines.append("")
    return "\n".join(lines)


# ================================================================ HTML =====
def to_html(reports, title="Web Accessibility Report"):
    def esc(s):
        return _html.escape(str(s))

    css = """
    body{font-family:-apple-system,Segoe UI,Roboto,Helvetica,Arial,sans-serif;
         max-width:1150px;margin:0 auto;padding:24px;color:#1b2430;background:#fafbfc;}
    h1{margin-bottom:4px;} .meta{color:#667;font-size:14px;margin-bottom:24px;}
    table{border-collapse:collapse;width:100%;margin-bottom:24px;background:#fff;}
    th,td{border:1px solid #dde2e8;padding:8px 10px;font-size:14px;text-align:left;}
    th{background:#eef1f5;}
    .pageblock{border:1px solid #dde2e8;border-radius:10px;padding:20px 24px;
              margin-bottom:32px;background:#fff;}
    .pagehead{display:flex;justify-content:space-between;align-items:baseline;
              flex-wrap:wrap;gap:8px;}
    .badge{display:inline-block;padding:2px 10px;border-radius:12px;font-size:12px;
           font-weight:600;color:#fff;}
    .row{border:1px solid #e3e7ec;border-radius:8px;padding:12px 16px;margin:10px 0;}
    .row h4{margin:0 0 6px 0;font-size:15px;}
    .rowmeta{font-size:12px;color:#556;margin-bottom:6px;}
    .issue{background:#f7f8fa;border-radius:6px;padding:8px 14px;margin-top:6px;font-size:13px;}
    .issue code{background:#eef1f5;padding:1px 5px;border-radius:4px;word-break:break-all;}
    .summarybar{display:flex;gap:10px;flex-wrap:wrap;margin:8px 0 16px 0;}
    .chip{border-radius:8px;padding:6px 12px;font-size:13px;font-weight:600;}
    """
    parts = [f"<!DOCTYPE html><html><head><meta charset='utf-8'>"
             f"<title>{esc(title)}</title><style>{css}</style></head><body>"]
    parts.append(f"<h1>{esc(title)}</h1>")
    overall = _overall(reports)
    parts.append(f"<div class='meta'>Generated {datetime.datetime.now():%Y-%m-%d %H:%M} "
                 f"&middot; {len(reports)} page(s) audited "
                 f"&middot; Overall conformance score: {esc(_score_str(overall['score_pct']))} "
                 f"({overall['pass']} pass / {overall['fail']} fail of {overall['determinable']} "
                 f"determinable checks; {overall['review']} flagged for manual review)</div>")

    if len(reports) > 1:
        parts.append("<table><tr><th>#</th><th>Page</th><th>Pass</th><th>Fail</th>"
                      "<th>Manual Review</th><th>Score</th></tr>")
        for i, r in enumerate(reports, start=1):
            s = r["summary"]
            parts.append(f"<tr><td>{i}</td><td><a href='{esc(r['url'])}'>{esc(r['display_name'])}</a></td>"
                         f"<td>{s['pass']}</td><td>{s['fail']}</td>"
                         f"<td>{s['review']}</td><td>{esc(_score_str(s['score_pct']))}</td></tr>")
        parts.append("</table>")

    for r in reports:
        s = r["summary"]
        parts.append("<div class='pageblock'>")
        parts.append(f"<div class='pagehead'><h2>{esc(r['display_name'])}</h2>"
                     f"<span class='badge' style='background:#2c3e50'>Score: {esc(_score_str(s['score_pct']))}</span></div>")
        parts.append(f"<div class='meta'>{esc(r['url'])}</div>")
        parts.append("<div class='summarybar'>"
                     f"<span class='chip' style='background:{STATUS_FILL['PASS']};color:{STATUS_COLORS['PASS']}'>Pass {s['pass']}</span>"
                     f"<span class='chip' style='background:{STATUS_FILL['FAIL']};color:{STATUS_COLORS['FAIL']}'>Fail {s['fail']}</span>"
                     f"<span class='chip' style='background:{STATUS_FILL['REVIEW']};color:{STATUS_COLORS['REVIEW']}'>Manual review {s['review']}</span>"
                     "</div>")
        if r.get("errors"):
            parts.append(f"<div class='meta'>Notes: {esc('; '.join(r['errors']))}</div>")
        for row in r["rows"]:
            color = STATUS_COLORS.get(row["status"], "555")
            parts.append("<div class='row'>")
            parts.append(f"<h4>{esc(row['id'])} &mdash; {esc(row['title'])} "
                         f"<span class='badge' style='background:{color}'>{esc(STATUS_LABEL[row['status']])}</span></h4>")
            parts.append(f"<div class='rowmeta'>VGAR ID: {esc(row['id'])} &middot; Category: {esc(row['category'])} &middot; "
                         f"Automation: {esc(TIER_LABEL[row['automation']])}</div>")
            parts.append(f"<div class='rowmeta'>WCAG: {esc(_wcag_str(row))}</div>")
            parts.append(f"<div class='rowmeta'>{esc(row['method'])}</div>")
            if row["status"] == "REVIEW" and row.get("procedure"):
                parts.append("<div class='issue'><strong>How to test manually</strong>"
                             + (f"<br><strong>Tool:</strong> {esc(row['tool'])}" if row.get("tool") else "")
                             + f"<br><strong>Steps:</strong> {esc(row['procedure'])}"
                             + (f"<div style='white-space:pre-line'><strong>Pass / fail:</strong> {esc(row['outcomes'])}</div>" if row.get("outcomes") else "")
                             + "</div>")
            for issue in row["issues"]:
                parts.append("<div class='issue'>"
                             f"<strong>VGAR:</strong> {esc(row['id'])} &middot; <strong>WCAG:</strong> {esc(_issue_wcag(issue, row))}<br>"
                             f"<strong>Location:</strong> <code>{esc(_issue_location(issue))}</code><br>"
                             f"<strong>Issue:</strong> {esc(issue['message'])}")
                if issue.get("html"):
                    parts.append(f"<br><strong>Code:</strong> <code>{esc(issue['html'][:250])}</code>")
                if issue.get("fix"):
                    parts.append(f"<br><strong>Fix:</strong> {esc(issue['fix'])}")
                parts.append("</div>")
            parts.append("</div>")
        parts.append("</div>")

    parts.append("</body></html>")
    return "\n".join(parts)


# ================================================================ XLSX =====
LINK_FONT = Font(color="0563C1", underline="single")


def _safe_sheet_name(name):
    bad = set('[]:*?/\\')
    cleaned = "".join(c for c in name if c not in bad)
    return cleaned[:31] or "Page"


def to_xlsx(reports, path, title="Web Accessibility Report"):
    wb = Workbook()

    used_names = set()
    sheet_names = []
    for r in reports:
        base = _safe_sheet_name(r["display_name"])
        name = base
        n = 2
        while name in used_names:
            name = _safe_sheet_name(f"{base[:28]}_{n}")
            n += 1
        used_names.add(name)
        sheet_names.append(name)

    overall = _overall(reports)

    # ---- Summary sheet ----
    ws = wb.active
    ws.title = "Summary"
    ws.append([title])
    ws["A1"].font = Font(size=16, bold=True)
    ws.append([f"Generated {datetime.datetime.now():%Y-%m-%d %H:%M}"])
    ws.append([f"{len(reports)} page(s) audited - Overall conformance score: {_score_str(overall['score_pct'])} "
               f"({overall['pass']} pass / {overall['fail']} fail of {overall['determinable']} determinable checks; "
               f"{overall['review']} flagged for manual review)"])
    ws["A3"].font = Font(size=12, bold=True, color="2C3E50")
    ws.append([])
    header_row = 5
    header = ["#", "Page (click to open its results tab)", "URL", "Pass", "Fail",
              "Manual Review", "Automated", "Partial", "Manual-only", "Score"]
    ws.append(header)
    for c in ws[header_row]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2C3E50")
    for r, sheet_name in zip(reports, sheet_names):
        s = r["summary"]
        ws.append([None, r["display_name"], r["url"], s["pass"], s["fail"], s["review"],
                   s["automated_tier"], s["partial_tier"], s["manual_tier"],
                   _score_str(s["score_pct"])])
        row_idx = ws.max_row
        ws.cell(row=row_idx, column=1, value=row_idx - header_row)
        cell = ws.cell(row=row_idx, column=2)
        cell.hyperlink = f"#'{sheet_name}'!A1"
        cell.font = LINK_FONT
    total_row_idx = ws.max_row + 1
    ws.append(["", "TOTAL (all pages in this batch)", "", overall["pass"], overall["fail"],
               overall["review"], "", "", "", _score_str(overall["score_pct"])])
    for c in ws[total_row_idx]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="EEF1F5")
    for col, width in zip("ABCDEFGHIJ", [4, 34, 40, 7, 7, 12, 10, 8, 11, 8]):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = f"A{header_row + 1}"

    # ---- One sheet per page (grouped findings, one row per issue) ----
    for r, sheet_name in zip(reports, sheet_names):
        ws2 = wb.create_sheet(sheet_name)
        ws2.append(["<< Back to Summary"])
        ws2.cell(row=1, column=1).hyperlink = "#'Summary'!A1"
        ws2.cell(row=1, column=1).font = LINK_FONT
        s = r["summary"]
        ws2.append([r["display_name"], r["url"], f"Conformance score: {_score_str(s['score_pct'])}"])
        ws2["A2"].font = Font(size=13, bold=True)
        ws2.append([])
        cols = ["VGAR ID", "Requirement", "WCAG", "Category", "Automation", "Status",
                "Location", "Code Snippet", "Issue / Finding", "Recommended Fix", "Method",
                "Manual Test Steps"]
        ws2.append(cols)
        header_row_idx = ws2.max_row
        for c in ws2[header_row_idx]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="2C3E50")
            c.alignment = Alignment(wrap_text=True, vertical="top")

        for row in r["rows"]:
            status_label = STATUS_LABEL[row["status"]]
            fill = STATUS_FILL.get(row["status"], "FFFFFF")
            steps = _manual_steps(row) if row["status"] == "REVIEW" else ""
            if row["issues"]:
                for issue in row["issues"]:
                    ws2.append([row["id"], row["title"], _issue_wcag(issue, row), row["category"],
                                TIER_LABEL[row["automation"]], status_label, _issue_location(issue),
                                issue.get("html", ""), issue["message"], issue.get("fix", ""),
                                row["method"], steps])
                    last_row = ws2.max_row
                    for c in ws2[last_row]:
                        c.alignment = Alignment(wrap_text=True, vertical="top")
                    ws2.cell(row=last_row, column=6).fill = PatternFill("solid", fgColor=fill)
            else:
                ws2.append([row["id"], row["title"], _wcag_str(row), row["category"],
                            TIER_LABEL[row["automation"]], status_label, "", "",
                            "" if row["status"] == "PASS" else row["requirement"],
                            "", row["method"], steps])
                last_row = ws2.max_row
                for c in ws2[last_row]:
                    c.alignment = Alignment(wrap_text=True, vertical="top")
                ws2.cell(row=last_row, column=6).fill = PatternFill("solid", fgColor=fill)

        widths = [10, 30, 26, 14, 16, 22, 30, 40, 40, 44, 40, 60]
        for i, w in enumerate(widths, start=1):
            ws2.column_dimensions[get_column_letter(i)].width = w
        ws2.freeze_panes = f"A{header_row_idx + 1}"

    wb.save(path)
    return path
